import os
import io
import pytest
import fitz  # PyMuPDF
import docx
from app import create_app
from app.extensions import db
from app.models.interview import Interview, InterviewStatus


@pytest.fixture
def app():
    """Create Flask application configured for testing."""
    os.environ["FLASK_ENV"] = "testing"
    os.environ["LLM_PROVIDER"] = "mock"
    os.environ["ELEVENLABS_AUTH_MODE"] = "signed_url"
    os.environ["ELEVENLABS_API_KEY"] = "mock-api-key"
    os.environ["ELEVENLABS_AGENT_ID"] = "mock-agent-id"

    test_app = create_app("testing")

    with test_app.app_context():
        db.create_all()
        yield test_app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    """Test client fixture with authenticated session."""
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["user_id"] = "test-user-123"
        sess["user_name"] = "Test User"
        sess["user_email"] = "test@example.com"
    return c


@pytest.fixture
def unauth_client(app):
    """Unauthenticated test client fixture."""
    return app.test_client()


@pytest.fixture
def sample_pdf_bytes():
    """Create a sample valid text-based PDF in memory."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        fitz.Point(50, 72),
        "John Doe\nContent Creator & Technical Writer\n"
        "Experience: 5 years creating multi-channel SEO blogs, tutorials, and brand guides.\n"
        "Skills: Copywriting, SEO keyword research, Contentful, Analytics, Editorial Calendars.\n"
        "Achievements: Increased organic traffic by 140% over 12 months."
    )
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_scanned_pdf_bytes():
    """Create a PDF with no text (simulates scanned document)."""
    doc = fitz.open()
    doc.new_page()  # Blank page without text layer
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Create a sample valid DOCX in memory."""
    doc = docx.Document()
    doc.add_heading("Job Description - Senior Content Strategist", level=1)
    doc.add_paragraph("Responsibilities: Plan, produce, and optimize high-impact editorial content across web and social channels.")
    doc.add_paragraph("Requirements: 3+ years writing experience, deep SEO knowledge, brand storytelling, and analytics tracking.")
    
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


@pytest.fixture
def prepared_interview(app):
    """Create an interview in 'prepared' status with plan and mock transcript."""
    interview = Interview(
        candidate_name="Alex Mercer",
        candidate_email="alex.mercer@example.com",
        job_title="Content Creator",
        jd_text="Looking for an experienced Content Creator to lead editorial strategy, SEO writing, and audience growth.",
        resume_text="Senior Content Writer with 4 years experience at TechPublish. Strong SEO, CMS, and storytelling expertise.",
        duration_minutes=15,
        difficulty="Intermediate",
        llm_provider="mock",
        voice_engine="elevenlabs",
        status=InterviewStatus.PREPARED,
        interview_plan={
            "job_title": "Content Creator",
            "required_competencies": [
                {
                    "id": "content_strategy",
                    "name": "Content Strategy",
                    "weight": 50,
                    "difficulty": "Intermediate",
                    "objective": "Assess editorial planning and channel mix",
                    "expected_evidence": "Concrete workflows and scheduling"
                },
                {
                    "id": "seo_writing",
                    "name": "SEO & Audience Research",
                    "weight": 50,
                    "difficulty": "Intermediate",
                    "objective": "Assess keyword targeting and intent matching",
                    "expected_evidence": "Search intent mapping and analytics"
                }
            ],
            "sections": []
        },
        transcript=[
            {
                "role": "interviewer",
                "message": "Welcome Alex! Could you tell me about your approach to building a quarterly content strategy?",
                "time_in_call_secs": 10
            },
            {
                "role": "candidate",
                "message": "I start by conducting thorough audience research and customer interview surveys. Then I map out topic clusters and keyword opportunities with high search intent.",
                "time_in_call_secs": 35
            },
            {
                "role": "interviewer",
                "message": "How do you evaluate whether a piece of content was successful after publication?",
                "time_in_call_secs": 65
            },
            {
                "role": "candidate",
                "message": "I monitor engagement metrics like average dwell time, scroll depth, organic signups, and social shares, iterating on headers and CTAs if drop-off is high.",
                "time_in_call_secs": 95
            }
        ]
    )
    db.session.add(interview)
    db.session.commit()
    return interview
