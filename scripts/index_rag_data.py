#!/usr/bin/env python3
"""
Turnkey Pinecone Vector DB Ingestion Script for Interview Questions RAG.
Reads 'Copy of Interviews.docx', normalizes and cleans the data,
generates dense vector embeddings, and upserts to Pinecone.

Usage:
    python scripts/index_rag_data.py
    python scripts/index_rag_data.py --file "Copy of Interviews.docx" --index "interview-questions"
"""

import os
import sys
import json
import re
import argparse
from pathlib import Path
from typing import List, Dict, Any

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# Ensure project root is in sys.path
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from dotenv import load_dotenv
load_dotenv()

try:
    import docx
except ImportError:
    print("Error: python-docx is required. Please install with: pip install python-docx")
    sys.exit(1)


def parse_interviews_docx(file_path: str) -> List[Dict[str, Any]]:
    """
    Parse 'Copy of Interviews.docx', normalize merged/empty role cells,
    split bundled multi-question cells into clean individual questions,
    and associate domain keywords from the job descriptions.
    """
    doc_path = Path(file_path)
    if not doc_path.exists():
        raise FileNotFoundError(f"Document not found at: {file_path}")

    doc = docx.Document(doc_path)

    # 1. Extract Job Descriptions
    jds = {}
    current_role = None
    current_lines = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if "Instructional Designer" in text and ("About" in text or "Role" in text):
            if current_role and current_lines:
                jds[current_role] = "\n".join(current_lines).strip()
            current_role = "Instructional Designer"
            current_lines = [text]
        elif "Graphic Designer" in text and ("About" in text or "Role" in text):
            if current_role and current_lines:
                jds[current_role] = "\n".join(current_lines).strip()
            current_role = "Graphic Designer"
            current_lines = [text]
        elif "Interior Designer" in text and ("Job" in text or "description" in text):
            if current_role and current_lines:
                jds[current_role] = "\n".join(current_lines).strip()
            current_role = "Interior Designer"
            current_lines = [text]
        elif "HR :" in text or "US IT Recruiter" in text:
            if current_role and current_lines:
                jds[current_role] = "\n".join(current_lines).strip()
            current_role = "HR / US IT Recruiter"
            current_lines = [text]
        else:
            if current_role:
                current_lines.append(text)

    if current_role and current_lines:
        jds[current_role] = "\n".join(current_lines).strip()

    print(f"✓ Extracted {len(jds)} Job Descriptions from document:")
    for role_name in jds:
        print(f"   • {role_name}")

    # 2. Extract Questions from Tables with normalization
    cleaned_records = []

    def clean_question_text(q_str: str) -> str:
        q = q_str.strip()
        # Clean numbering prefixes e.g. "1.", "2)", "3:"
        q = re.sub(r'^\d+[\.:\)]\s*', '', q).strip()
        # Fix known typos from raw doc
        q = q.replace("WHat", "What")
        if q == "Do you have experience PowerPoint." or q == "Do you have experience with PowerPoint?":
            q = "Do you have experience with PowerPoint and corporate presentation design?"
        if q == "Project I was proud of?":
            q = "Can you describe a flagship project you led that you are particularly proud of?"
        if q.lower().startswith("what is the depth of wardrobe"):
            q = "What is the standard functional depth of a residential wardrobe and how do you customize internal configurations?"
        if "What are the core materials for kitchen cabinets" in q and not q.endswith("?"):
            q += "?"
        if not q.endswith("?") and not q.endswith("."):
            q += "?"
        return q[:1].upper() + q[1:]

    current_table_role = None
    q_counter = 1

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows[1:], 1):
            cell_role = row.cells[0].text.strip()
            comp = row.cells[1].text.strip()
            raw_qs = row.cells[2].text.strip()

            if cell_role:
                current_table_role = cell_role

            if not comp and not raw_qs:
                continue

            role = current_table_role or "General"

            # Split bundled lines & numbered items
            lines = raw_qs.split('\n')
            for line in lines:
                parts = re.split(r'(?:^|\s+)\d+[\.:\)]\s*', line.strip())
                for p_text in parts:
                    cleaned_q = clean_question_text(p_text)
                    if len(cleaned_q) > 8:
                        # Category heuristic
                        cat = "Core Assessment"
                        q_lower = cleaned_q.lower()
                        if "portfolio" in q_lower or "project" in q_lower:
                            cat = "Project & Portfolio Deep-Dive"
                        elif "tool" in q_lower or "software" in q_lower or "powerpoint" in q_lower or "autocad" in q_lower or "sketchup" in q_lower:
                            cat = "Technical & Tool Proficiency"
                        elif "deadline" in q_lower or "time" in q_lower:
                            cat = "Deadline & Time Management"
                        elif "color" in q_lower or "typography" in q_lower or "layout" in q_lower:
                            cat = "Design Foundations"
                        elif "collaborat" in q_lower or "team" in q_lower:
                            cat = "Cross-Functional Collaboration"
                        elif "wirefram" in q_lower or "ux" in q_lower:
                            cat = "UX & Wireframing"
                        elif "kitchen" in q_lower or "wardrobe" in q_lower or "material" in q_lower:
                            cat = "Technical Specifications & Materials"
                        elif "online" in q_lower or "blended" in q_lower or "model" in q_lower:
                            cat = "Learning Strategy & Pedagogy"

                        record = {
                            "id": f"{role.lower().replace(' ', '-')[:15]}-{comp.lower().replace(' ', '-')[:12]}-{q_counter}",
                            "role": role,
                            "company": comp or "Top Tier Employer",
                            "category": cat,
                            "question": cleaned_q,
                            "jd_keywords": jds.get(role, "")[:400]
                        }
                        cleaned_records.append(record)
                        q_counter += 1

    # 3. Add HR / US IT Recruiter Questions (Table 3 in doc was blank)
    hr_questions = [
        ("Top IT Staffing", "Sourcing & Boolean Search", "How do you construct advanced Boolean search strings and utilize LinkedIn Recruiter and job boards to source passive niche technical talent?"),
        ("Top IT Staffing", "US Tax Terms & Compliance", "Can you explain the structural and legal differences between W2, 1099, and C2C (Corp-to-Corp) employment models in US recruitment?"),
        ("Top IT Staffing", "Candidate Assessment & Screening", "Walk me through your initial screening call. How do you assess candidate technical fit, salary expectations, and cultural alignment?"),
        ("Top IT Staffing", "Offer Management & Retention", "How do you handle a situation where an offered candidate receives a counteroffer or has competing offers?"),
        ("Top IT Staffing", "ATS & Pipeline Metrics", "What applicant tracking systems (ATS) have you utilized, and how do you track submittal-to-placement conversion rates?")
    ]

    for comp, cat, q in hr_questions:
        cleaned_records.append({
            "id": f"hr-recruiter-{q_counter}",
            "role": "HR / US IT Recruiter",
            "company": comp,
            "category": cat,
            "question": q,
            "jd_keywords": jds.get("HR / US IT Recruiter", "")[:400]
        })
        q_counter += 1

    print(f"✓ Extracted and normalized {len(cleaned_records)} high-signal interview questions.")
    return cleaned_records


def generate_embedding(text: str, model: str = "text-embedding-3-small") -> List[float]:
    """Generate dense embedding vector for record."""
    openai_key = os.getenv("OPENAI_API_KEY") or os.getenv("LLM_API_KEY", "")
    if openai_key and not openai_key.startswith("your_") and not openai_key.startswith("AIza"):
        try:
            from openai import OpenAI
            client = OpenAI(api_key=openai_key)
            resp = client.embeddings.create(model=model, input=text[:4000])
            return resp.data[0].embedding
        except Exception as e:
            print(f"OpenAI embedding error: {e}")

    # Fallback to Gemini
    gemini_key = os.getenv("GEMINI_API_KEY") or (openai_key if openai_key.startswith("AIza") else "")
    if gemini_key and not gemini_key.startswith("your_"):
        try:
            from google import genai
            g_client = genai.Client(api_key=gemini_key)
            res = g_client.models.embed_content(model="text-embedding-004", contents=text[:4000])
            if res and res.embedding:
                return list(res.embedding.values)
        except Exception as e:
            print(f"Gemini embedding error: {e}")

    raise RuntimeError("No valid OpenAI or Gemini API key found to generate vector embeddings.")


def index_to_pinecone(
    records: List[Dict[str, Any]],
    index_name: str = "interview-questions",
    namespace: str = "interviews",
    cloud: str = "aws",
    region: str = "us-east-1"
):
    """Create Pinecone index if needed, generate embeddings, and upsert records."""
    pinecone_key = os.getenv("PINECONE_API_KEY", "").strip()
    if not pinecone_key or pinecone_key.startswith("your_"):
        print("\n⚠️ PINECONE_API_KEY is not set in your .env file.")
        print("Please add your key from https://app.pinecone.io to .env:")
        print("   PINECONE_API_KEY=pcsk_...\n")
        return False

    try:
        from pinecone import Pinecone, ServerlessSpec
        pc = Pinecone(api_key=pinecone_key)

        # 1. Determine embedding dimension
        test_vec = generate_embedding("sample test prompt")
        dim = len(test_vec)
        print(f"✓ Verified embedding vector dimension: {dim}")

        # 2. Check / Create Index
        existing_indexes = [idx.name for idx in pc.list_indexes()]
        if index_name not in existing_indexes:
            print(f"Creating serverless Pinecone index '{index_name}' (dim={dim}, metric=cosine)...")
            pc.create_index(
                name=index_name,
                dimension=dim,
                metric="cosine",
                spec=ServerlessSpec(cloud=cloud, region=region)
            )
            print(f"✓ Index '{index_name}' successfully created.")
        else:
            print(f"✓ Using existing Pinecone index: '{index_name}'")

        idx = pc.Index(index_name)

        # 3. Embed and Upsert Vectors
        vectors_to_upsert = []
        print(f"\nEmbedding and preparing {len(records)} questions for Pinecone...")

        for r in records:
            # Dense text representation combining role, company, focus, and question
            embed_text = (
                f"Role: {r['role']}\n"
                f"Company: {r['company']}\n"
                f"Focus: {r['category']}\n"
                f"Question: {r['question']}\n"
                f"Keywords: {r.get('jd_keywords', '')}"
            )
            vec = generate_embedding(embed_text)
            metadata = {
                "role": r["role"],
                "company": r["company"],
                "category": r["category"],
                "question": r["question"]
            }
            vectors_to_upsert.append({
                "id": r["id"],
                "values": vec,
                "metadata": metadata
            })

        print(f"Upserting {len(vectors_to_upsert)} vectors to Pinecone (namespace='{namespace}')...")
        idx.upsert(vectors=vectors_to_upsert, namespace=namespace)
        print(f"✓ Successfully upserted {len(vectors_to_upsert)} questions to Pinecone!")

        # 4. Verification Test Query
        test_role = "Instructional Designer"
        print(f"\n--- Verifying Search: '{test_role}' ---")
        q_vec = generate_embedding(f"Target Role: {test_role}\nInterview Questions")
        search_res = idx.query(vector=q_vec, top_k=3, include_metadata=True, namespace=namespace)
        for m in getattr(search_res, "matches", []):
            print(f" • [{m.metadata.get('company')}] {m.metadata.get('question')} (Score: {m.score:.3f})")

        return True

    except Exception as e:
        print(f"Error indexing to Pinecone: {e}")
        return False


def main():
    parser = argparse.ArgumentParser(description="Index Copy of Interviews.docx into Pinecone RAG")
    parser.add_argument("--file", default="Copy of Interviews.docx", help="Path to input docx file")
    parser.add_argument("--index", default=os.getenv("PINECONE_INDEX", "interview-questions"), help="Pinecone index name")
    parser.add_argument("--namespace", default=os.getenv("PINECONE_NAMESPACE", "interviews"), help="Pinecone namespace")
    parser.add_argument("--export-json", default="interview_knowledge_base.json", help="Path to export clean JSON")
    args = parser.parse_args()

    print("==============================================================")
    print("Pinecone Vector DB Knowledge Base Indexer")
    print("==============================================================")

    # 1. Parse and normalize
    records = parse_interviews_docx(args.file)

    # 2. Export clean JSON for user inspection
    export_path = ROOT_DIR / args.export_json
    with open(export_path, "w", encoding="utf-8") as f:
        json.dump(records, f, indent=2, ensure_ascii=False)
    print(f"✓ Exported clean normalized knowledge base to: {export_path}")

    # 3. Index to Pinecone if key is present
    success = index_to_pinecone(
        records,
        index_name=args.index,
        namespace=args.namespace
    )

    if success:
        print("\n🎉 Pinecone RAG is fully loaded and ready for live interviews!")
    else:
        print("\nNote: Local normalized knowledge base JSON is saved. Once you set PINECONE_API_KEY in .env, run this script again to push vectors to Pinecone.")


if __name__ == "__main__":
    main()
