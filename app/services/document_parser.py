import os
import re
from pathlib import Path
from typing import Optional, Tuple
from pydantic import BaseModel, Field
import fitz  # PyMuPDF
import docx
from werkzeug.utils import secure_filename


class DocumentParsingError(Exception):
    """Raised when document extraction fails or file is invalid."""
    pass


class ScannedDocumentError(DocumentParsingError):
    """Raised when PDF has no text layer (scanned or image-only)."""
    pass


# Structured Pydantic models for extracted information
class JDAnalysis(BaseModel):
    job_title: str = Field(..., description="Target job title extracted or inferred from JD")
    responsibilities: list[str] = Field(default_factory=list, description="Primary duties and responsibilities")
    required_skills: list[str] = Field(default_factory=list, description="Mandatory technical or domain skills")
    preferred_skills: list[str] = Field(default_factory=list, description="Bonus or nice-to-have skills")
    experience_requirements: list[str] = Field(default_factory=list, description="Years or depth of experience required")
    tools_and_technologies: list[str] = Field(default_factory=list, description="Software tools, platforms, frameworks")
    role_specific_competencies: list[str] = Field(default_factory=list, description="Key competencies evaluated")
    evaluation_criteria: list[str] = Field(default_factory=list, description="What success looks like in this role")


class ResumeAnalysis(BaseModel):
    skills: list[str] = Field(default_factory=list, description="Skills explicitly evidenced in the resume")
    work_experience: list[str] = Field(default_factory=list, description="Summary of work history and roles")
    projects: list[str] = Field(default_factory=list, description="Key projects completed")
    education: list[str] = Field(default_factory=list, description="Educational background")
    certifications: list[str] = Field(default_factory=list, description="Relevant certifications or licenses")
    achievements: list[str] = Field(default_factory=list, description="Quantified achievements and outcomes")
    portfolio_links: list[str] = Field(default_factory=list, description="Web links, GitHub, or portfolio URLs if present")


def normalize_extracted_text(text: str) -> str:
    """Normalize extracted text by cleaning up whitespace and control characters."""
    if not text:
        return ""
    # Replace non-breaking spaces and form feeds
    text = text.replace("\xa0", " ").replace("\x0c", "\n")
    # Normalize multiple newlines and carriage returns
    text = re.sub(r"\r\n|\r", "\n", text)
    # Condense consecutive horizontal whitespace
    text = re.sub(r"[ \t]+", " ", text)
    # Condense 3+ newlines to 2
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


class DocumentParserService:
    """Handles parsing, validation, and extraction of text from PDF, DOCX, and TXT files."""

    ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB

    @classmethod
    def allowed_file(cls, filename: str) -> bool:
        return "." in filename and filename.rsplit(".", 1)[1].lower() in cls.ALLOWED_EXTENSIONS

    @classmethod
    def validate_file(cls, file_storage, max_size: int = MAX_FILE_SIZE) -> Tuple[bool, Optional[str]]:
        """Validate filename, extension, and content size."""
        if not file_storage or not file_storage.filename:
            return False, "No file was selected."
        
        filename = secure_filename(file_storage.filename)
        if not filename:
            return False, "Invalid filename."
            
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in cls.ALLOWED_EXTENSIONS:
            return False, f"Unsupported file extension '.{ext}'. Supported formats are: PDF, DOCX, TXT."
        
        # Check size by seeking to end
        file_storage.seek(0, os.SEEK_END)
        size = file_storage.tell()
        file_storage.seek(0)
        
        if size > max_size:
            max_mb = max_size // (1024 * 1024)
            return False, f"File exceeds the maximum allowable size of {max_mb} MB."
            
        if size == 0:
            return False, "Uploaded file is empty (0 bytes)."
            
        return True, None

    @classmethod
    def extract_from_pdf(cls, file_stream_or_path) -> str:
        """Extract text from PDF using PyMuPDF (fitz) and detect scanned/empty documents."""
        try:
            if isinstance(file_stream_or_path, (str, Path)):
                doc = fitz.open(file_stream_or_path)
            else:
                file_bytes = file_stream_or_path.read()
                doc = fitz.open(stream=file_bytes, filetype="pdf")

            extracted_chunks = []
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text("text")
                if text:
                    extracted_chunks.append(text)
            doc.close()

            full_text = "\n".join(extracted_chunks)
            normalized = normalize_extracted_text(full_text)
            
            # Check if text is non-existent or suspiciously small
            if len(normalized.replace(" ", "")) < 25:
                raise ScannedDocumentError(
                    "The uploaded PDF contains no readable text. It appears to be scanned or image-based. "
                    "Please provide a text-searchable PDF or convert it using OCR."
                )
            return normalized

        except ScannedDocumentError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Failed to process PDF document: {str(e)}")

    @classmethod
    def extract_from_docx(cls, file_stream_or_path) -> str:
        """Extract text from DOCX document."""
        try:
            doc = docx.Document(file_stream_or_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text inside tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)
            normalized = normalize_extracted_text(full_text)
            if not normalized:
                raise DocumentParsingError("The uploaded DOCX file does not contain any readable text.")
            return normalized

        except Exception as e:
            if isinstance(e, DocumentParsingError):
                raise
            raise DocumentParsingError(f"Failed to process DOCX document: {str(e)}")

    @classmethod
    def extract_from_txt(cls, file_stream_or_path) -> str:
        """Extract text from plain text file."""
        try:
            if isinstance(file_stream_or_path, (str, Path)):
                with open(file_stream_or_path, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
            else:
                content = file_stream_or_path.read().decode("utf-8", errors="replace")

            normalized = normalize_extracted_text(content)
            if not normalized:
                raise DocumentParsingError("The uploaded text file is empty.")
            return normalized

        except Exception as e:
            if isinstance(e, DocumentParsingError):
                raise
            raise DocumentParsingError(f"Failed to process text file: {str(e)}")

    @classmethod
    def extract_text(cls, file_storage, filename: Optional[str] = None) -> str:
        """Dispatcher to extract text based on filename extension."""
        fname = filename or getattr(file_storage, "filename", "")
        if not fname:
            raise DocumentParsingError("No filename provided for text extraction.")

        ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
        if ext == "pdf":
            return cls.extract_from_pdf(file_storage)
        elif ext == "docx":
            return cls.extract_from_docx(file_storage)
        elif ext == "txt":
            return cls.extract_from_txt(file_storage)
        else:
            raise DocumentParsingError(f"Unsupported file format: '.{ext}'")
